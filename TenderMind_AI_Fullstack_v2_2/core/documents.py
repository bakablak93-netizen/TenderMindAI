"""Извлечение текста без передачи документов сторонним сервисам, кроме выбранного AI-провайдера."""
import csv
import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_FILE_MB = 15
MAX_TEXT_CHARS = 55000


def extract_text(name: str, content: bytes) -> str:
    if len(content) > MAX_FILE_MB * 1024 * 1024:
        raise ValueError(f"Максимальный размер файла — {MAX_FILE_MB} МБ")
    ext = Path(name).suffix.lower()
    if ext == '.pdf':
        pdf = PdfReader(io.BytesIO(content))
        if pdf.is_encrypted:
            raise ValueError('PDF с паролем: разблокируйте документ перед загрузкой')
        text = '\n\n'.join((page.extract_text() or '') for page in pdf.pages[:100])
        if not text.strip():
            raise ValueError('PDF не содержит извлекаемого текста. Для сканирования нужен OCR (пока не поддерживается).')
    elif ext == '.docx':
        doc = Document(io.BytesIO(content))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(chunks)
    elif ext == '.xlsx':
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets[:10]:
            chunks.append(f'Лист: {ws.title}')
            for row in ws.iter_rows(max_row=500, max_col=30, values_only=True):
                parts = [str(v) if v is not None else '' for v in row]
                if any(parts):
                    chunks.append(' | '.join(parts))
        wb.close()
        text = '\n'.join(chunks)
    elif ext in ('.csv', '.txt'):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('cp1251', errors='replace')
    else:
        raise ValueError('Поддерживаются только PDF, DOCX, XLSX, CSV и TXT')
    return text[:MAX_TEXT_CHARS].strip()
